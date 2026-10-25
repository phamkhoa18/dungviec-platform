from __future__ import annotations

import base64
import json
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "filebaocao"
ASSETS = OUT / "report_assets"
DIAGRAMS = ROOT / "docs" / "diagrams"
NOTEBOOK = ROOT / "kqtrain" / "TRAIN_CV_JD_FIT_MODEL.ipynb"
EXCEL_100 = OUT / "DATASET_TEST_100_SAMPLES_RESULTS.xlsx"
ASSETS.mkdir(exist_ok=True)

REPORT2 = OUT / "BaoCaoLan2_DacTaVaThietKe_HoanChinh.docx"
REPORT3 = OUT / "BaoCaoLan3_HuanLuyenAI_DeHieu.docx"

NAVY = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F3F8FC"
GRAY = "D9D9D9"
PALE_GRAY = "F5F5F5"
WHITE = "FFFFFF"
BLACK = "000000"
FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"


def extract_notebook_image(cell_index: int, output_index: int, target: Path) -> None:
    nb = json.loads(NOTEBOOK.read_text(encoding="utf-8"))
    payload = nb["cells"][cell_index]["outputs"][output_index]["data"]["image/png"]
    if isinstance(payload, list):
        payload = "".join(payload)
    target.write_bytes(base64.b64decode(payload))


def pil_font(size: int, bold: bool = False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT, size)


def line_chart(target: Path, title: str, series: list[tuple[str, list[float], str]], y_min: float, y_max: float) -> None:
    w, h = 1500, 700
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    left, top, right, bottom = 120, 100, 1430, 590
    d.text((w / 2, 35), title, font=pil_font(34, True), fill="#1F2937", anchor="ma")
    for i in range(6):
        y = top + (bottom - top) * i / 5
        val = y_max - (y_max - y_min) * i / 5
        d.line((left, y, right, y), fill="#E5E7EB", width=2)
        d.text((left - 18, y), f"{val:.2f}", font=pil_font(20), fill="#4B5563", anchor="rm")
    for e in range(1, 6):
        x = left + (right - left) * (e - 1) / 4
        d.text((x, bottom + 22), str(e), font=pil_font(21), fill="#374151", anchor="ma")
    d.text((w / 2, bottom + 70), "Epoch", font=pil_font(23), fill="#374151", anchor="ma")
    for name, vals, color in series:
        pts = []
        for i, v in enumerate(vals):
            x = left + (right - left) * i / 4
            y = bottom - (v - y_min) / (y_max - y_min) * (bottom - top)
            pts.append((x, y))
        d.line(pts, fill=color, width=6)
        for x, y in pts:
            d.ellipse((x - 8, y - 8, x + 8, y + 8), fill=color)
    lx = left + 20
    for name, _, color in series:
        d.rectangle((lx, 625, lx + 35, 650), fill=color)
        d.text((lx + 48, 637), name, font=pil_font(20), fill="#111827", anchor="lm")
        lx += 430
    img.save(target)


def grouped_bar_chart(target: Path) -> None:
    w, h = 1500, 760
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((w / 2, 35), "So sánh kết quả trên tập test 1.759 mẫu", font=pil_font(34, True), fill="#1F2937", anchor="ma")
    left, top, right, bottom = 130, 110, 1430, 610
    for i in range(8):
        y = bottom - (bottom - top) * i / 7
        val = i * 10
        d.line((left, y, right, y), fill="#E5E7EB", width=2)
        d.text((left - 18, y), str(val), font=pil_font(20), fill="#4B5563", anchor="rm")
    names = ["TF-IDF + LR", "Mô hình tham chiếu", "DistilBERT đề xuất"]
    acc = [45.59, 58.40, 44.86]
    f1 = [42.66, 43.20, 37.79]
    centers = [360, 780, 1200]
    colors = ["#4C78A8", "#F58518"]
    for j, center in enumerate(centers):
        for k, value in enumerate((acc[j], f1[j])):
            x1 = center - 90 + k * 100
            x2 = x1 + 80
            y1 = bottom - value / 70 * (bottom - top)
            d.rectangle((x1, y1, x2, bottom), fill=colors[k])
            d.text(((x1 + x2) / 2, y1 - 12), f"{value:.2f}", font=pil_font(19), fill="#111827", anchor="ms")
        d.text((center, bottom + 28), names[j], font=pil_font(20), fill="#111827", anchor="ma")
    d.rectangle((480, 690, 515, 715), fill=colors[0])
    d.text((530, 702), "Accuracy (%)", font=pil_font(20), fill="#111827", anchor="lm")
    d.rectangle((800, 690, 835, 715), fill=colors[1])
    d.text((850, 702), "Macro F1 × 100", font=pil_font(20), fill="#111827", anchor="lm")
    img.save(target)


def sample_summary_chart(target: Path, notebook_correct: int, excel_correct: int) -> None:
    w, h = 1500, 700
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((375, 45), "Kết quả notebook gốc", font=pil_font(32, True), fill="#1F2937", anchor="ma")
    d.ellipse((145, 140, 605, 600), fill="#E45756")
    angle = int(360 * notebook_correct / 100)
    d.pieslice((145, 140, 605, 600), -90, -90 + angle, fill="#54A24B")
    d.ellipse((255, 250, 495, 490), fill="white")
    d.text((375, 350), f"{notebook_correct}%", font=pil_font(52, True), fill="#111827", anchor="mm")
    d.text((375, 415), "dự đoán đúng", font=pil_font(24), fill="#374151", anchor="mm")
    d.text((1080, 45), "Đối chiếu hai tệp minh chứng", font=pil_font(32, True), fill="#1F2937", anchor="ma")
    base, max_h = 590, 420
    values = [notebook_correct, excel_correct]
    names = ["Notebook", "Excel hiện có"]
    colors = ["#4C78A8", "#F58518"]
    for i, value in enumerate(values):
        cx = 930 + i * 300
        bh = max_h * value / 50
        d.rectangle((cx - 85, base - bh, cx + 85, base), fill=colors[i])
        d.text((cx, base - bh - 18), f"{value}%", font=pil_font(26, True), fill="#111827", anchor="ms")
        d.text((cx, base + 22), names[i], font=pil_font(21), fill="#111827", anchor="ma")
    d.text((1080, 640), "Hai tệp cần được tái xuất từ cùng một lần chạy", font=pil_font(19), fill="#6B7280", anchor="ma")
    img.save(target)


def equation_image(target: Path) -> None:
    w, h = 1800, 220
    img = Image.new("RGBA", (w, h), (255, 255, 255, 0))
    d = ImageDraw.Draw(img)
    text = "S = 100 × [P(Good Fit) + 0,5 × P(Potential Fit)]"
    d.text((w / 2, h / 2), text, font=pil_font(54), fill="#111111", anchor="mm")
    img.save(target)


def generate_assets() -> None:
    extract_notebook_image(6, 2, ASSETS / "phan_bo_du_lieu.png")
    extract_notebook_image(14, 3, ASSETS / "confusion_matrix_test.png")

    epochs = np.arange(1, 6)
    train_loss = [1.976347, 1.591700, 1.479288, 1.300854, 1.174633]
    val_loss = [0.856049, 0.776578, 0.723612, 0.692373, 0.683805]
    val_acc = [0.532051, 0.584402, 0.584402, 0.658120, 0.660256]
    macro_f1 = [0.537395, 0.586235, 0.588008, 0.657637, 0.660856]
    line_chart(ASSETS / "training_loss.png", "Loss theo epoch", [("Training loss", train_loss, "#4C78A8"), ("Validation loss", val_loss, "#F58518")], 0.5, 2.1)
    line_chart(ASSETS / "training_metrics.png", "Chỉ số validation theo epoch", [("Validation accuracy", val_acc, "#54A24B"), ("Validation Macro F1", macro_f1, "#B279A2")], 0.45, 0.72)
    canvas = Image.new("RGB", (1500, 1420), "white")
    canvas.paste(Image.open(ASSETS / "training_loss.png"), (0, 0))
    canvas.paste(Image.open(ASSETS / "training_metrics.png"), (0, 710))
    canvas.save(ASSETS / "training_curves.png")
    grouped_bar_chart(ASSETS / "benchmark_test.png")

    wb = load_workbook(EXCEL_100, read_only=True, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    truth_i = headers.index("Ground Truth Label")
    pred_i = headers.index("AI Predicted Label")
    truths, preds = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        truths.append(str(row[truth_i]))
        preds.append(str(row[pred_i]))
    labels = ["No Fit", "Potential Fit", "Good Fit"]
    correct = sum(a == b for a, b in zip(truths, preds))
    pred_counts = [preds.count(x) for x in labels]
    sample_summary_chart(ASSETS / "test_100_summary.png", 31, correct)
    equation_image(ASSETS / "fit_score_equation.png")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_font(run, size=12, bold=False, italic=False, color=BLACK) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 13), ("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 13)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_font(run, size=10)


def add_cover(doc: Document, report_no: int, subtitle: str, date: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(f"BÁO CÁO CÔNG VIỆC LẦN {report_no}")
    set_font(r, 13, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(subtitle)
    set_font(r, 13, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run("Đề tài: Xây dựng website hỗ trợ tuyển dụng ứng dụng AI trong phân tích CV và đánh giá mức độ phù hợp với vị trí tuyển dụng")
    set_font(r, 13)
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    table.columns[0].width = Inches(2.25)
    table.columns[1].width = Inches(4.15)
    entries = [
        ("Sinh viên thực hiện:", "Phạm Đăng Khoa - 24810114"),
        ("Giảng viên hướng dẫn:", "ThS. Đoàn Thanh Phong"),
        ("Thời gian báo cáo:", date),
    ]
    for row, (label, value) in zip(table.rows, entries):
        prevent_row_split(row)
        row.cells[0].text = label
        row.cells[1].text = value
        for i, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, 150, 150, 150, 150)
            for rr in cell.paragraphs[0].runs:
                set_font(rr, 13, bold=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.75)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_together = True
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, 13, True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r, 13)
    else:
        r = p.add_run(text)
        set_font(r, 13)


def add_bullets(doc: Document, items: list[str], numbered: bool = False) -> None:
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph(style="List Bullet" if not numbered else None)
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(3)
        text = f"{index}. {item}" if numbered else item
        r = p.add_run(text)
        set_font(r, 13)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None, font_size=10.5) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = text
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in cell.paragraphs[0].runs:
            set_font(r, font_size, True)
    for ridx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, text in enumerate(values):
            cell = row.cells[i]
            cell.text = str(text)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.25
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    set_font(r, font_size)
    if widths:
        total = sum(widths)
        usable = 6.55
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(usable * w / total)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.45) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.add_run().add_picture(str(path), width=Inches(width))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(caption)
    set_font(r, 10, italic=True)


def add_use_case(doc: Document, code: str, title: str, actor: str, pre: str, main: list[str], exceptions: list[str]) -> None:
    add_heading(doc, f"{code} {title}", 2)
    add_body(doc, f"Tác nhân: {actor}. Điều kiện tiên quyết: {pre}")
    p = doc.add_paragraph()
    r = p.add_run("Luồng chính")
    set_font(r, 12, True)
    add_bullets(doc, main, numbered=True)
    p = doc.add_paragraph()
    r = p.add_run("Ngoại lệ")
    set_font(r, 12, True)
    add_bullets(doc, exceptions)


def build_report2() -> None:
    doc = Document()
    configure_doc(doc)
    add_cover(doc, 2, "ĐẶC TẢ YÊU CẦU VÀ THIẾT KẾ HỆ THỐNG", "23/10/2026")

    add_heading(doc, "1 Mục tiêu và phạm vi báo cáo", 1)
    add_body(doc, "Báo cáo lần 2 chuyển kết quả khảo sát ở lần 1 thành đặc tả có thể dùng để thiết kế và lập trình. Nội dung tập trung vào tác nhân, yêu cầu chức năng, các luồng nghiệp vụ, cấu trúc dữ liệu và kiến trúc dự kiến. Báo cáo này chưa trình bày số liệu huấn luyện hoặc kết quả đánh giá mô hình AI; toàn bộ nội dung thực nghiệm được tách sang Báo cáo lần 3.")
    add_bullets(doc, [
        "Xác định hai nhóm người dùng chính gồm ứng viên và nhà tuyển dụng.",
        "Mô tả các chức năng bao phủ chu trình đăng tin, nộp hồ sơ, phân tích AI và sàng lọc ứng viên.",
        "Thiết kế cơ sở dữ liệu PostgreSQL gồm 06 bảng và các ràng buộc toàn vẹn.",
        "Đề xuất kiến trúc Next.js, FastAPI, AI Engine, PostgreSQL và lưu trữ tệp CV.",
    ])

    add_heading(doc, "2 Tác nhân và phạm vi chức năng", 1)
    add_table(doc, ["Tác nhân", "Trách nhiệm", "Quyền chính"], [
        ["Ứng viên", "Tìm việc, quản lý CV, nộp hồ sơ và theo dõi tiến trình.", "Chỉ xem và thay đổi dữ liệu thuộc tài khoản của mình."],
        ["Nhà tuyển dụng", "Quản lý công ty, đăng tin, xem ứng viên, cập nhật trạng thái.", "Chỉ quản lý tin và hồ sơ ứng tuyển thuộc công ty mình."],
        ["Dịch vụ AI", "Đọc nội dung CV, trích xuất thông tin và chấm điểm CV-JD.", "Chạy tự động sau khi đơn ứng tuyển được tạo; không phải tài khoản người dùng."],
    ], [1.2, 2.6, 2.8])
    add_body(doc, "Phạm vi hiện tại không bao gồm cổng quản trị viên, thanh toán, phỏng vấn trực tuyến hoặc gửi email hàng loạt. Cách giới hạn này giữ quy mô phù hợp với một tiểu luận chuyên ngành và tập trung vào giá trị chính là hỗ trợ sàng lọc CV bằng AI.")
    add_figure(doc, DIAGRAMS / "1_use_case.png", "Hình 2.1 Sơ đồ Use Case tổng quan", 6.55)

    add_heading(doc, "3 Yêu cầu chức năng", 1)
    add_table(doc, ["Mã", "Yêu cầu", "Kết quả mong đợi"], [
        ["FR-01", "Đăng ký và đăng nhập", "Xác thực tài khoản theo vai trò candidate hoặc hr."],
        ["FR-02", "Tìm kiếm việc làm", "Tìm theo từ khóa và lọc theo địa điểm, mức lương hoặc ngành nghề."],
        ["FR-03", "Quản lý CV", "Tải lên PDF hoặc DOCX tối đa 5 MB và chọn CV dùng để ứng tuyển."],
        ["FR-04", "Nộp hồ sơ", "Tạo một đơn ứng tuyển duy nhất cho mỗi cặp ứng viên và tin tuyển dụng."],
        ["FR-05", "Theo dõi hồ sơ", "Ứng viên xem trạng thái và thời điểm cập nhật gần nhất."],
        ["FR-06", "Quản lý công ty", "HR cập nhật thông tin doanh nghiệp đại diện."],
        ["FR-07", "Quản lý tin tuyển dụng", "HR tạo, sửa, đóng hoặc mở tin trong phạm vi công ty."],
        ["FR-08", "Phân tích CV", "Hệ thống trích xuất văn bản và các thông tin cần thiết từ CV."],
        ["FR-09", "Chấm điểm CV-JD", "Hệ thống trả nhãn, xác suất và điểm phù hợp từ 0 đến 100."],
        ["FR-10", "Xếp hạng ứng viên", "HR xem danh sách được sắp theo điểm phù hợp giảm dần."],
        ["FR-11", "Quản lý tiến trình", "HR cập nhật trạng thái và ghi chú nội bộ cho từng hồ sơ."],
    ], [0.7, 2.0, 4.0], 10)

    add_heading(doc, "4 Đặc tả các Use Case cốt lõi", 1)
    add_use_case(doc, "UC-04", "Nộp hồ sơ ứng tuyển", "Ứng viên", "Ứng viên đã đăng nhập, tin còn hiệu lực và có CV hợp lệ.", [
        "Ứng viên mở chi tiết tin tuyển dụng và chọn ứng tuyển.",
        "Ứng viên chọn CV, nhập thư giới thiệu và xác nhận gửi.",
        "Hệ thống kiểm tra loại tệp, dung lượng, thời hạn tin và hồ sơ trùng.",
        "Hệ thống tạo application ở trạng thái pending và kích hoạt tác vụ AI nền.",
        "Ứng viên nhận thông báo nộp hồ sơ thành công.",
    ], [
        "Tệp không phải PDF hoặc DOCX, vượt 5 MB hoặc không đọc được.",
        "Ứng viên đã nộp vào cùng vị trí hoặc tin đã đóng/hết hạn.",
        "AI gặp lỗi: đơn vẫn được lưu và chờ xử lý lại, không làm mất hồ sơ.",
    ])
    add_use_case(doc, "UC-07", "Đăng và quản lý tin tuyển dụng", "Nhà tuyển dụng", "HR đã đăng nhập và có hồ sơ công ty.", [
        "HR nhập tiêu đề, ngành nghề, địa điểm, lương, nội dung JD, kỹ năng và hạn nộp.",
        "Hệ thống kiểm tra dữ liệu và lưu ở trạng thái draft hoặc active.",
        "HR có thể chỉnh sửa nội dung hoặc chuyển tin sang closed.",
    ], ["Thiếu trường bắt buộc, ngày hết hạn không hợp lệ hoặc HR không sở hữu công ty."])
    add_use_case(doc, "UC-08", "Xem bảng xếp hạng ứng viên", "Nhà tuyển dụng", "Tin tuyển dụng có hồ sơ và HR có quyền truy cập.", [
        "Hệ thống lấy danh sách application của tin.",
        "Kết quả được nối với ai_evaluations và sắp xếp theo match_score giảm dần.",
        "HR xem nhãn, xác suất, kỹ năng khớp, kỹ năng thiếu và CV gốc.",
    ], ["Kết quả AI chưa có hoặc thất bại được hiển thị ở trạng thái chờ, không tự động loại ứng viên."])
    add_use_case(doc, "UC-09", "Cập nhật trạng thái hồ sơ", "Nhà tuyển dụng", "HR sở hữu tin tuyển dụng tương ứng.", [
        "HR mở hồ sơ và chọn trạng thái mới.",
        "Hệ thống kiểm tra quyền, lưu trạng thái, ghi chú và updated_at.",
        "Ứng viên nhìn thấy trạng thái mới trong trang hồ sơ của mình.",
    ], ["Nếu cập nhật thất bại, hệ thống giữ trạng thái cũ và thông báo để HR thử lại."])

    add_heading(doc, "5 Quy trình nghiệp vụ", 1)
    add_body(doc, "Quy trình bắt đầu khi HR tạo tin tuyển dụng ở trạng thái active. Ứng viên tìm việc, tải CV và xác nhận ứng tuyển. Backend lưu đơn trước, sau đó mới gọi tác vụ AI để tránh mất dữ liệu khi mô hình gặp lỗi. Khi kết quả AI được ghi vào cơ sở dữ liệu, HR có thể dùng điểm số để ưu tiên đọc hồ sơ nhưng quyết định cuối cùng vẫn thuộc về HR.")
    add_figure(doc, DIAGRAMS / "2_quy_trinh_nghiep_vu.png", "Hình 2.2 Quy trình nghiệp vụ từ đăng tin đến sàng lọc", 6.45)
    add_figure(doc, DIAGRAMS / "6_trang_thai_ho_so.png", "Hình 2.3 Vòng đời đơn ứng tuyển", 6.45)

    add_heading(doc, "6 Thiết kế cơ sở dữ liệu", 1)
    add_body(doc, "Cơ sở dữ liệu dùng PostgreSQL và UUID cho toàn bộ khóa chính, khóa ngoại. Bảng applications giữ vai trò trung tâm vì liên kết người dùng, tin tuyển dụng, CV đã nộp và kết quả AI. Quan hệ application với ai_evaluations là 1 đến 0 hoặc 1: một đơn có thể chưa có kết quả trong thời gian tác vụ AI đang chạy.")
    add_figure(doc, DIAGRAMS / "3_erd.png", "Hình 2.4 Sơ đồ ERD gồm 06 bảng", 6.55)
    schemas = [
        ("users", "id, email, password_hash, role, full_name, phone, avatar_url, created_at, updated_at", "email duy nhất; role thuộc candidate hoặc hr"),
        ("companies", "id, hr_user_id, company_name, logo_url, website, address, description, created_at", "hr_user_id tham chiếu users"),
        ("jobs", "id, company_id, title, job_category, location, salary_range, experience_required, jd_text, required_skills, deadline, status", "status thuộc draft, active, closed"),
        ("cv_profiles", "id, user_id, cv_title, file_path, file_name, file_type, raw_text, extracted_skills, years_of_experience, education_level", "file_type thuộc pdf hoặc docx"),
        ("applications", "id, job_id, candidate_id, cv_profile_id, cover_letter, status, hr_notes, applied_at, updated_at", "duy nhất theo job_id và candidate_id"),
        ("ai_evaluations", "id, application_id, match_score, predicted_label, ba xác suất, matched_skills, missing_skills, processed_time_ms, created_at", "application_id duy nhất; match_score từ 0 đến 100"),
    ]
    add_table(doc, ["Bảng", "Trường dữ liệu chính", "Ràng buộc quan trọng"], [list(x) for x in schemas], [1.2, 3.7, 2.0], 9.5)

    add_heading(doc, "7 Kiến trúc hệ thống dự kiến", 1)
    add_figure(doc, DIAGRAMS / "5_kien_truc.png", "Hình 2.5 Kiến trúc Next.js, FastAPI, AI Engine và PostgreSQL", 6.35)
    add_bullets(doc, [
        "Frontend Next.js cung cấp giao diện riêng cho ứng viên và HR.",
        "Backend FastAPI xử lý xác thực JWT, nghiệp vụ tuyển dụng, kiểm tra quyền và điều phối tác vụ AI.",
        "AI Engine gồm bộ đọc tài liệu, mô-đun trích xuất thông tin CV và mô-đun đánh giá mức phù hợp CV-JD.",
        "PostgreSQL lưu dữ liệu có cấu trúc; file storage lưu CV gốc và checkpoint mô hình.",
    ])

    add_heading(doc, "8 Thiết kế đầu vào và đầu ra của mô đun AI", 1)
    add_body(doc, "Ở Báo cáo lần 2, AI chỉ được mô tả ở mức giao diện tích hợp. Đầu vào gồm văn bản CV, nội dung JD và mã application. Đầu ra dự kiến gồm nhãn No Fit, Potential Fit hoặc Good Fit; xác suất của ba lớp; điểm S; kỹ năng khớp và kỹ năng còn thiếu. Các con số đánh giá mô hình không thuộc phạm vi báo cáo này.")
    add_figure(doc, ASSETS / "fit_score_equation.png", "Công thức điểm phù hợp dùng để xếp hạng ứng viên", 4.9)
    add_body(doc, "Điểm S là giá trị hỗ trợ sắp xếp, không thay thế quyết định tuyển dụng. Khi kết quả AI chưa có hoặc không đáng tin cậy, HR vẫn phải xem CV gốc và đưa ra quyết định thủ công.")

    add_heading(doc, "9 Quy tắc nghiệp vụ và xử lý ngoại lệ", 1)
    add_table(doc, ["Mã", "Quy tắc"], [
        ["BR-01", "Chỉ chấp nhận CV PDF hoặc DOCX, dung lượng tối đa 5 MB."],
        ["BR-02", "Mỗi ứng viên chỉ có một application cho một job."],
        ["BR-03", "Tin draft hoặc closed không nhận hồ sơ; tin quá deadline tự chuyển closed."],
        ["BR-04", "HR chỉ xem và cập nhật hồ sơ thuộc tin của công ty mình."],
        ["BR-05", "Tổng ba xác suất AI phải xấp xỉ 1 và điểm S nằm trong khoảng 0 đến 100."],
        ["BR-06", "Lỗi AI không làm thất bại thao tác nộp đơn; hệ thống ghi nhận và xử lý lại."],
        ["BR-07", "AI chỉ hỗ trợ ưu tiên hồ sơ; không tự động từ chối ứng viên."],
    ], [0.8, 5.9])

    add_heading(doc, "10 Kết quả cần xác nhận và kế hoạch tiếp theo", 1)
    add_body(doc, "Sau Báo cáo lần 2, ba nội dung cần được giảng viên xác nhận là phạm vi hai vai trò người dùng, lược đồ 06 bảng và công thức điểm S. Công việc của Báo cáo lần 3 là trình bày quá trình huấn luyện DistilBERT, kết quả validation, kết quả test độc lập, đánh giá xếp hạng và các hạn chế quan sát được từ dữ liệu thực nghiệm.")
    doc.save(REPORT2)


def build_report3() -> None:
    doc = Document()
    configure_doc(doc)
    add_cover(doc, 3, "HUẤN LUYỆN VÀ ĐÁNH GIÁ MÔ HÌNH AI", "30/10/2026")

    add_heading(doc, "1 Mục tiêu và phạm vi thực nghiệm", 1)
    add_body(doc, "Báo cáo lần 3 trình bày toàn bộ quá trình huấn luyện mô hình DistilBERT Cross-Encoder dùng để đánh giá mức độ phù hợp giữa CV và mô tả công việc. Số liệu được lấy trực tiếp từ notebook đã chạy trên Google Colab và tập kết quả 100 mẫu. Mô đun BERT NER mới dừng ở thiết kế kiến trúc, chưa có kết quả huấn luyện trong notebook này nên không được báo cáo như một mô hình đã hoàn thành.")
    add_bullets(doc, [
        "Xây dựng baseline TF-IDF kết hợp Logistic Regression để có mốc so sánh.",
        "Fine-tune distilbert-base-uncased với đầu vào ghép cặp CV-JD và loss có trọng số lớp.",
        "Đánh giá phân loại trên validation, test độc lập và mẫu kiểm tra phân tầng 100 cặp.",
        "Đánh giá chất lượng xếp hạng bằng NDCG@5 và NDCG@10.",
    ])

    add_heading(doc, "2 Môi trường thực nghiệm", 1)
    add_table(doc, ["Thành phần", "Cấu hình"], [
        ["Nền tảng", "Google Colab"],
        ["GPU", "NVIDIA Tesla T4, VRAM 15,64 GB"],
        ["PyTorch", "2.11.0+cu130"],
        ["Thư viện chính", "Transformers, Datasets, PyTorch, scikit-learn, pandas, NumPy"],
        ["Seed", "42"],
        ["Checkpoint nền", "distilbert-base-uncased, 6 tầng Transformer, khoảng 66 triệu tham số"],
    ], [1.7, 5.0])

    add_heading(doc, "3 Dữ liệu và tiền xử lý", 1)
    add_body(doc, "Tập cnamuangtoun/resume-job-description-fit gồm 8.000 cặp CV-JD tiếng Anh, được chia sẵn thành 6.241 mẫu train và 1.759 mẫu test. Nhãn gồm No Fit, Potential Fit và Good Fit. Phân bố lớp không cân bằng: No Fit chiếm xấp xỉ một nửa dữ liệu, hai lớp còn lại chiếm khoảng một phần tư mỗi lớp.")
    add_table(doc, ["Tập", "No Fit", "Potential Fit", "Good Fit", "Tổng"], [
        ["Train gốc", "3.143", "1.556", "1.542", "6.241"],
        ["Test", "857", "444", "458", "1.759"],
    ], [1.2, 1.2, 1.5, 1.3, 1.0])
    add_figure(doc, ASSETS / "phan_bo_du_lieu.png", "Hình 3.1 Phân bố nhãn của tập train và test", 6.1)
    add_body(doc, "Tiền xử lý dùng MD5 của cặp CV-JD để loại bản ghi trùng, loại văn bản ngắn dưới 50 ký tự và kiểm tra nhãn mâu thuẫn. Train giảm từ 6.241 xuống 6.234 mẫu do loại 7 cặp trùng; test giữ nguyên 1.759 mẫu. Sau đó train sạch được chia phân tầng 85% cho huấn luyện và 15% cho validation.")
    add_table(doc, ["Tập sau xử lý", "Số mẫu", "Mục đích"], [
        ["Train", "5.298", "Cập nhật trọng số mô hình"],
        ["Validation", "936", "Chọn checkpoint theo Macro F1"],
        ["Test độc lập", "1.759", "Đánh giá tổng quát hóa cuối cùng"],
    ], [2.0, 1.2, 3.5])

    add_heading(doc, "4 Baseline TF IDF và Logistic Regression", 1)
    add_body(doc, "Baseline dùng hai TF-IDF vectorizer độc lập: tối đa 5.000 đặc trưng cho CV và 3.000 đặc trưng cho JD, n-gram từ 1 đến 2, sau đó ghép hai vector và huấn luyện Logistic Regression với class_weight='balanced'. Baseline đạt Accuracy 45,59% và Macro F1 0,4266 trên 1.759 mẫu test.")
    add_table(doc, ["Lớp", "Precision", "Recall", "F1", "Support"], [
        ["No Fit", "0,6120", "0,5484", "0,5785", "857"],
        ["Potential Fit", "0,2951", "0,4347", "0,3515", "444"],
        ["Good Fit", "0,4125", "0,3035", "0,3497", "458"],
    ], [1.7, 1.2, 1.2, 1.2, 1.0])

    add_heading(doc, "5 Mô hình DistilBERT Cross Encoder", 1)
    add_body(doc, "Mỗi mẫu được biểu diễn theo cấu trúc [CLS] CV [SEP] JD [SEP]. Để không làm mất toàn bộ JD khi CV dài, bộ dữ liệu cấp tối đa 384 token cho CV và 125 token cho JD; ba token đặc biệt đưa tổng chiều dài tối đa về đúng 512 token.")
    add_body(doc, "Mất cân bằng lớp được xử lý bằng Cross-Entropy có trọng số nghịch đảo tần suất. Trọng số thực tế là 0,6614 cho No Fit, 1,3372 cho Potential Fit và 1,3511 cho Good Fit.")
    add_table(doc, ["Siêu tham số", "Giá trị"], [
        ["Learning rate", "2e-5"],
        ["Batch train / eval", "8 / 16"],
        ["Gradient accumulation", "2"],
        ["Số epoch", "5"],
        ["Weight decay", "0,01"],
        ["Warmup steps", "100"],
        ["Mixed precision", "FP16 khi có CUDA"],
        ["Tiêu chí chọn mô hình", "Macro F1 trên validation"],
        ["Scheduler", "Cosine"],
    ], [2.4, 4.3])
    add_figure(doc, DIAGRAMS / "4_ai_pipeline.png", "Hình 3.2 Pipeline huấn luyện và suy luận AI", 6.45)

    add_heading(doc, "6 Kết quả huấn luyện trên validation", 1)
    rows = [
        ["1", "1,9763", "0,8560", "53,21%", "0,5374", "0,4527", "0,5042", "0,6552"],
        ["2", "1,5917", "0,7766", "58,44%", "0,5862", "0,5178", "0,5680", "0,6729"],
        ["3", "1,4793", "0,7236", "58,44%", "0,5880", "0,4713", "0,5974", "0,6954"],
        ["4", "1,3009", "0,6924", "65,81%", "0,6576", "0,6490", "0,6192", "0,7046"],
        ["5", "1,1746", "0,6838", "66,03%", "0,6609", "0,6422", "0,6329", "0,7075"],
    ]
    add_table(doc, ["Epoch", "Train loss", "Val loss", "Accuracy", "Macro F1", "F1 No Fit", "F1 Potential", "F1 Good Fit"], rows, [0.7, 1.0, 1.0, 1.0, 1.0, 1.0, 1.1, 1.0], 8.5)
    add_figure(doc, ASSETS / "training_curves.png", "Hình 3.3 Loss và chỉ số validation qua 5 epoch", 6.25)
    add_body(doc, "Validation loss giảm liên tục từ 0,8560 xuống 0,6838; Accuracy và Macro F1 tăng mạnh ở epoch 4 và đạt cao nhất ở epoch 5. Checkpoint epoch 5 được chọn theo Macro F1 validation. Tuy nhiên, validation tốt không đủ để kết luận mô hình tổng quát hóa tốt; kết quả test độc lập phải được xem xét riêng.")

    add_heading(doc, "7 Kết quả trên tập test độc lập", 1)
    add_table(doc, ["Mô hình", "Accuracy", "Macro F1", "F1 No Fit", "F1 Potential", "F1 Good Fit"], [
        ["TF-IDF + Logistic Regression", "45,59%", "0,4266", "0,5785", "0,3515", "0,3497"],
        ["Mô hình tham chiếu domynom", "58,40%", "0,4320", "0,6900", "0,2800", "0,3260"],
        ["DistilBERT đề xuất", "44,86%", "0,3779", "0,6015", "0,2099", "0,3224"],
    ], [2.5, 1.0, 1.0, 1.0, 1.1, 1.0], 9)
    add_figure(doc, ASSETS / "benchmark_test.png", "Hình 3.4 So sánh Accuracy và Macro F1 trên tập test", 6.2)
    add_figure(doc, ASSETS / "confusion_matrix_test.png", "Hình 3.5 Ma trận nhầm lẫn của baseline và DistilBERT", 6.45)
    add_body(doc, "Trên test độc lập, DistilBERT đạt Accuracy 44,86% và Macro F1 0,3779, thấp hơn baseline lần lượt 0,73 điểm phần trăm và 0,0487 Macro F1. Điểm yếu lớn nhất là Potential Fit với F1 0,2099. Chênh lệch lớn giữa validation và test cho thấy mô hình chưa tổng quát hóa ổn định; có thể tồn tại lệch phân phối giữa các tập hoặc mô hình học các tín hiệu đặc thù của train/validation.")
    add_body(doc, "Vì vậy, checkpoint hiện tại được xem là kết quả thực nghiệm để tiếp tục cải thiện, chưa đủ cơ sở để triển khai như bộ lọc tự động. Nếu tích hợp vào website, kết quả chỉ nên dùng để sắp xếp tham khảo và luôn cho phép HR xem CV gốc.")

    add_heading(doc, "8 Đánh giá chất lượng xếp hạng", 1)
    add_figure(doc, ASSETS / "fit_score_equation.png", "Công thức chuyển xác suất thành điểm phù hợp S", 4.9)
    add_table(doc, ["Chỉ số", "Kết quả", "Cách hiểu"], [
        ["Mean NDCG@5", "0,5842", "Đánh giá thứ tự 5 ứng viên đầu trong mỗi nhóm mô phỏng."],
        ["Mean NDCG@10", "0,7256", "Đánh giá thứ tự tối đa 10 ứng viên trong mỗi nhóm mô phỏng."],
    ], [1.6, 1.2, 3.9])
    add_body(doc, "NDCG được tính bằng cách gán ngẫu nhiên các mẫu test vào các nhóm công việc mô phỏng, không phải các nhóm ứng viên thực tế cùng ứng tuyển một JD. Kết quả cho thấy điểm S có tiềm năng dùng cho xếp hạng, nhưng chưa thể dùng để khẳng định hiệu quả tuyển dụng ngoài thực tế. Cần xây dựng bộ đánh giá theo từng JD thật để kiểm chứng thêm.")

    add_heading(doc, "9 Kiểm thử định tính", 1)
    add_table(doc, ["Kịch bản", "Nhãn dự đoán", "Điểm S", "Nhận xét"], [
        ["Senior Python Backend đến Python Developer", "Good Fit", "71,1", "Phù hợp với kỳ vọng."],
        ["Civil Engineer đến Python Developer", "No Fit", "27,5", "Phù hợp với kỳ vọng."],
        ["Junior Data Analyst đến Senior Backend Developer", "Good Fit", "82,4", "Dự đoán quá cao; cho thấy mô hình nhạy với từ khóa Python và SQL."],
    ], [2.6, 1.2, 0.8, 2.3], 9.5)
    add_body(doc, "Kịch bản thứ ba là phản ví dụ quan trọng: hồ sơ chỉ có một năm kinh nghiệm và thiếu FastAPI, PostgreSQL, microservices nhưng mô hình vẫn dự đoán Good Fit với điểm 82,4. Điều này củng cố kết luận rằng mô hình chưa nên tự động ra quyết định tuyển dụng.")

    add_heading(doc, "10 Kiểm thử phân tầng 100 mẫu", 1)
    add_body(doc, "Tập kiểm tra bổ sung lấy ngẫu nhiên theo seed 42 gồm 34 No Fit, 33 Potential Fit và 33 Good Fit. Output được lưu trong notebook ghi nhận 31 dự đoán đúng trên 100 mẫu, tương ứng Accuracy 31%. Vì tập này được cân bằng lớp và nhỏ hơn test đầy đủ, con số chỉ dùng làm kiểm tra bổ sung, không thay thế kết quả trên 1.759 mẫu.")
    add_table(doc, ["Nguồn minh chứng", "Số mẫu đúng", "Accuracy", "Cách xử lý trong báo cáo"], [
        ["Output notebook đã chạy", "31/100", "31%", "Dùng làm kết quả chính vì gắn trực tiếp với phiên train được báo cáo."],
        ["File Excel hiện có", "39/100", "39%", "Ghi nhận là tệp từ lần xuất khác; cần tái xuất để đồng bộ."],
    ], [1.8, 1.2, 1.0, 2.8], 9.5)
    add_figure(doc, ASSETS / "test_100_summary.png", "Hình 3.6 Đối chiếu kết quả kiểm thử 100 mẫu giữa notebook và Excel", 6.1)

    doc.add_page_break()
    add_heading(doc, "11 Hạn chế và kế hoạch cải thiện", 1)
    add_table(doc, ["Hạn chế quan sát được", "Hướng cải thiện"], [
        ["Validation cao nhưng test giảm mạnh", "Kiểm tra trùng lặp theo resume/JD riêng, đánh giá theo nhóm nguồn dữ liệu và dùng cross-validation."],
        ["F1 Potential Fit chỉ đạt 0,2099", "Rà soát tiêu chí gán nhãn, thử ordinal classification hoặc gộp thành bài toán nhị phân trước."],
        ["Mô hình dự đoán quá cao khi có từ khóa liên quan", "Bổ sung negative samples khó và các đặc trưng định lượng về số năm kinh nghiệm, kỹ năng bắt buộc."],
        ["CV và JD dài bị cắt", "Thử mô hình hỗ trợ văn bản dài, chunking có pooling hoặc trích xuất phần kinh nghiệm/kỹ năng trước khi so khớp."],
        ["NDCG dùng nhóm công việc mô phỏng", "Tạo bộ dữ liệu theo từng JD với nhiều ứng viên và thứ hạng do HR đánh giá."],
        ["BERT NER chưa có kết quả train", "Huấn luyện và đánh giá NER riêng bằng precision, recall và entity-level F1 trước khi công bố kết quả."],
    ], [3.0, 3.7], 10)

    add_heading(doc, "12 Kết luận", 1)
    add_body(doc, "Quy trình huấn luyện đã chạy hoàn chỉnh trên Tesla T4, lưu được checkpoint DistilBERT và tạo đủ kết quả validation, test, NDCG, kiểm thử định tính và kiểm thử 100 mẫu. Kết quả validation tốt nhất là Accuracy 66,03% và Macro F1 0,6609, nhưng test độc lập chỉ đạt Accuracy 44,86% và Macro F1 0,3779. Do đó, đóng góp hiện tại nằm ở việc xây dựng pipeline thực nghiệm và phát hiện rõ vấn đề tổng quát hóa; mô hình cần được cải thiện trước khi dùng trong môi trường tuyển dụng thật.")

    add_heading(doc, "13 Tài liệu và tệp minh chứng", 1)
    add_bullets(doc, [
        "Dataset: https://huggingface.co/datasets/cnamuangtoun/resume-job-description-fit",
        "Checkpoint nền: https://huggingface.co/distilbert/distilbert-base-uncased",
        "Notebook thực nghiệm: kqtrain/TRAIN_CV_JD_FIT_MODEL.ipynb",
        "Kết quả 100 mẫu: filebaocao/DATASET_TEST_100_SAMPLES_RESULTS.xlsx",
    ])
    doc.save(REPORT3)


def build_report2() -> None:
    """Báo cáo lần 2: trình bày yêu cầu và thiết kế, chưa đưa kết quả train."""
    doc = Document()
    configure_doc(doc)
    add_cover(doc, 2, "PHÂN TÍCH YÊU CẦU VÀ THIẾT KẾ HỆ THỐNG", "23/10/2026")

    add_heading(doc, "1. Mục tiêu công việc lần 2")
    add_body(doc, "Lần 2 tập trung làm rõ hệ thống cần có những chức năng nào, người dùng sử dụng ra sao và dữ liệu được lưu như thế nào. Kết quả huấn luyện AI chưa trình bày ở báo cáo này mà được đưa sang Báo cáo lần 3.")

    add_heading(doc, "2. Người dùng và nhu cầu sử dụng")
    add_table(doc, ["Đối tượng", "Nhu cầu sử dụng", "Giới hạn"], [
        ["Ứng viên", "Tạo tài khoản, quản lý CV, tìm việc, nộp hồ sơ và xem trạng thái xử lý.", "Chỉ xem và sửa dữ liệu của chính mình."],
        ["Nhà tuyển dụng", "Quản lý công ty, đăng tin, xem hồ sơ, xem điểm phù hợp và cập nhật kết quả tuyển dụng.", "Chỉ quản lý tin và hồ sơ thuộc công ty của mình."],
        ["Mô-đun AI", "Đọc nội dung CV và mô tả công việc, sau đó trả về nhãn và điểm phù hợp.", "Chỉ hỗ trợ sắp xếp hồ sơ, không tự quyết định tuyển hay loại."],
    ], [1.4, 3.4, 2.0], 11)
    add_figure(doc, DIAGRAMS / "1_use_case.png", "Hình 2.1. Sơ đồ chức năng tổng quát", 6.35)

    add_heading(doc, "3. Các chức năng chính của hệ thống")
    add_table(doc, ["Mã", "Chức năng", "Mô tả"], [
        ["CN-01", "Đăng ký, đăng nhập", "Xác thực tài khoản và phân biệt ứng viên với nhà tuyển dụng."],
        ["CN-02", "Quản lý CV", "Tải CV PDF hoặc DOCX, lưu và chọn CV khi ứng tuyển."],
        ["CN-03", "Tìm việc", "Tìm theo từ khóa, địa điểm, ngành nghề và trạng thái tin."],
        ["CN-04", "Nộp hồ sơ", "Gửi CV vào tin còn hạn và tạo hồ sơ ứng tuyển."],
        ["CN-05", "Theo dõi hồ sơ", "Ứng viên xem trạng thái hồ sơ đã nộp."],
        ["CN-06", "Quản lý công ty", "Nhà tuyển dụng cập nhật thông tin doanh nghiệp."],
        ["CN-07", "Quản lý tin tuyển dụng", "Tạo, sửa, mở hoặc đóng tin tuyển dụng."],
        ["CN-08", "Phân tích bằng AI", "So sánh nội dung CV với yêu cầu công việc và tạo điểm phù hợp."],
        ["CN-09", "Xem danh sách ứng viên", "Sắp xếp hồ sơ theo điểm AI để nhà tuyển dụng tham khảo."],
        ["CN-10", "Cập nhật kết quả", "Chuyển trạng thái sang đang xem, phù hợp, nhận hoặc từ chối."],
    ], [0.8, 2.0, 4.0], 10.5)

    add_heading(doc, "4. Quy trình hoạt động chính")
    add_table(doc, ["Bước", "Người thực hiện", "Nội dung"], [
        ["1", "Nhà tuyển dụng", "Đăng nhập và tạo tin tuyển dụng."],
        ["2", "Ứng viên", "Tìm tin phù hợp, chọn CV và nộp hồ sơ."],
        ["3", "Hệ thống", "Kiểm tra loại tệp, hạn nộp và hồ sơ đã nộp trước đó."],
        ["4", "Hệ thống", "Lưu hồ sơ trước, sau đó mới gửi dữ liệu sang mô-đun AI."],
        ["5", "Mô-đun AI", "Đọc CV, so sánh với mô tả công việc và trả về kết quả."],
        ["6", "Nhà tuyển dụng", "Xem danh sách đã sắp xếp, mở CV gốc và đánh giá."],
        ["7", "Nhà tuyển dụng", "Cập nhật trạng thái; ứng viên xem được kết quả mới."],
    ], [0.7, 1.8, 4.3], 10.5)
    add_figure(doc, DIAGRAMS / "2_quy_trinh_nghiep_vu.png", "Hình 2.2. Quy trình từ đăng tin đến xử lý hồ sơ", 6.2)

    add_heading(doc, "5. Thiết kế dữ liệu")
    add_body(doc, "Dữ liệu được lưu bằng MongoDB và chia thành 6 collection. Mỗi document có trường _id kiểu ObjectId. Các trường như user_id, company_id và job_id dùng để tham chiếu document ở collection khác. Collection applications là phần trung tâm, dùng để nối ứng viên, tin tuyển dụng, CV đã nộp và kết quả AI.")
    add_figure(doc, DIAGRAMS / "3_erd.png", "Hình 2.3. Sơ đồ dữ liệu MongoDB", 6.3)
    add_table(doc, ["Collection", "Dữ liệu được lưu"], [
        ["users", "Tài khoản, họ tên, email và vai trò người dùng."],
        ["companies", "Thông tin công ty của nhà tuyển dụng."],
        ["jobs", "Nội dung tin, yêu cầu, kỹ năng, hạn nộp và trạng thái."],
        ["cv_profiles", "Tệp CV, nội dung đã đọc và thông tin rút ra từ CV."],
        ["applications", "Ứng viên đã nộp vào tin nào và trạng thái hiện tại."],
        ["ai_evaluations", "Nhãn dự đoán, xác suất, điểm phù hợp và kỹ năng còn thiếu."],
    ], [1.7, 5.1], 11)

    add_heading(doc, "6. Công nghệ dự kiến sử dụng")
    add_table(doc, ["Công nghệ", "Dùng để làm gì"], [
        ["Next.js", "Xây dựng giao diện web cho ứng viên và nhà tuyển dụng."],
        ["FastAPI", "Xử lý đăng nhập, tin tuyển dụng, hồ sơ và kết nối mô-đun AI."],
        ["MongoDB", "Lưu dữ liệu dạng document cho tài khoản, công ty, tin tuyển dụng, hồ sơ và kết quả AI."],
        ["PDF/DOCX reader", "Đọc văn bản từ tệp CV trước khi đưa vào mô hình."],
        ["DistilBERT", "Đọc đồng thời CV và mô tả công việc để dự đoán mức phù hợp."],
        ["Lưu trữ tệp", "Lưu CV gốc để nhà tuyển dụng có thể mở và kiểm tra lại."],
    ], [2.0, 4.8], 11)
    add_figure(doc, DIAGRAMS / "5_kien_truc.png", "Hình 2.4. Kiến trúc dự kiến của hệ thống", 6.15)

    add_heading(doc, "7. Quy tắc cần áp dụng")
    add_table(doc, ["Mã", "Quy tắc"], [
        ["QT-01", "Chỉ nhận CV PDF hoặc DOCX, dung lượng tối đa 5 MB."],
        ["QT-02", "Một ứng viên chỉ nộp một hồ sơ cho một tin tuyển dụng."],
        ["QT-03", "Tin đã đóng hoặc hết hạn không được nhận thêm hồ sơ."],
        ["QT-04", "Nhà tuyển dụng chỉ xem hồ sơ thuộc tin của công ty mình."],
        ["QT-05", "Nếu AI bị lỗi, hồ sơ vẫn phải được lưu để xử lý lại."],
        ["QT-06", "Điểm AI chỉ dùng tham khảo; quyết định cuối cùng thuộc nhà tuyển dụng."],
    ], [1.0, 5.8], 11)

    add_heading(doc, "8. Kết quả công việc và bước tiếp theo")
    add_body(doc, "Sau lần 2, hệ thống đã có danh sách chức năng, quy trình chính, thiết kế 6 collection MongoDB và kiến trúc dự kiến. Bước tiếp theo là huấn luyện mô hình DistilBERT, kiểm tra kết quả trên dữ liệu chưa dùng để train và xác định những điểm cần cải thiện. Nội dung này được trình bày trong Báo cáo lần 3.")
    doc.save(REPORT2)


def build_report3() -> None:
    """Báo cáo lần 3: trình bày quá trình train và kiểm tra bằng cách dễ hiểu."""
    doc = Document()
    configure_doc(doc)
    add_cover(doc, 3, "HUẤN LUYỆN VÀ KIỂM TRA MÔ HÌNH AI", "30/10/2026")

    add_heading(doc, "1. Công việc thực hiện trong lần 3")
    add_body(doc, "Trong lần 3, em huấn luyện mô hình DistilBERT để so sánh một CV với một mô tả công việc (JD). Mô hình trả về một trong ba kết quả: không phù hợp, có thể phù hợp hoặc phù hợp.")
    add_body(doc, "Phần kiểm tra tập trung trả lời ba câu hỏi: mô hình đoán đúng được bao nhiêu mẫu, mô hình có nhận ra từng mức phù hợp hay không và kết quả có đủ ổn định để đưa vào website hay chưa. Phần BERT NER dùng để lấy thông tin chi tiết từ CV chưa được train trong notebook hiện có, vì vậy báo cáo này chưa ghi kết quả của phần đó.")

    add_heading(doc, "2. Công cụ đã sử dụng")
    add_table(doc, ["Công cụ", "Dùng để làm gì"], [
        ["Google Colab", "Chạy notebook train và lưu kết quả."],
        ["NVIDIA Tesla T4", "GPU dùng để train mô hình, bộ nhớ 15,64 GB."],
        ["Python và PyTorch 2.11", "Đọc dữ liệu, train và chạy thử mô hình."],
        ["Hugging Face Transformers", "Tải mô hình DistilBERT có sẵn để train tiếp với dữ liệu CV và JD."],
        ["scikit-learn", "Tính tỷ lệ đoán đúng và tạo mô hình đơn giản để so sánh."],
        ["distilbert-base-uncased", "Mô hình tiếng Anh ban đầu được chọn vì dữ liệu CV và JD là tiếng Anh."],
    ], [2.2, 4.6], 11)

    add_heading(doc, "3. Dữ liệu được dùng như thế nào")
    add_body(doc, "Dữ liệu lấy từ bộ cnamuangtoun/resume-job-description-fit trên Hugging Face. Mỗi mẫu gồm nội dung CV, nội dung JD và đáp án đúng. Đáp án có ba mức: No Fit là không phù hợp, Potential Fit là có thể phù hợp và Good Fit là phù hợp.")
    add_table(doc, ["Tập dữ liệu", "No Fit", "Potential Fit", "Good Fit", "Tổng"], [
        ["Train gốc", "3.143", "1.556", "1.542", "6.241"],
        ["Test", "857", "444", "458", "1.759"],
    ], [1.5, 1.2, 1.5, 1.3, 1.0], 10.5)
    add_figure(doc, ASSETS / "phan_bo_du_lieu.png", "Hình 3.1. Phân bố nhãn của dữ liệu", 5.9)
    add_body(doc, "Sau khi loại 7 cặp CV-JD bị trùng, dữ liệu còn lại được chia theo mục đích rõ ràng như sau.")
    add_table(doc, ["Phần dữ liệu", "Số mẫu", "Mục đích"], [
        ["Train", "5.298", "Cho mô hình đọc và học cách phân biệt ba mức phù hợp."],
        ["Validation", "936", "Kiểm tra sau mỗi epoch và chọn phiên bản mô hình tốt nhất."],
        ["Test", "1.759", "Chỉ dùng sau khi train xong để xem mô hình xử lý dữ liệu mới ra sao."],
    ], [1.5, 1.0, 4.3], 10.5)
    add_body(doc, "Tập test không được dùng để mô hình học hoặc chọn checkpoint. Vì vậy, kết quả trên 1.759 mẫu test được xem là kết quả kiểm tra cuối cùng.")

    add_heading(doc, "4. Các bước train mô hình")
    add_table(doc, ["Bước", "Công việc thực hiện"], [
        ["1", "Loại các cặp CV-JD bị trùng và kiểm tra dữ liệu quá ngắn hoặc sai nhãn."],
        ["2", "Đưa CV và JD vào cùng một mẫu để mô hình đọc và so sánh."],
        ["3", "Giới hạn CV ở 384 token và JD ở 125 token để tổng đầu vào không vượt quá 512 token."],
        ["4", "Tăng mức chú ý cho hai nhãn ít mẫu hơn để nhãn No Fit không lấn át kết quả."],
        ["5", "Cho mô hình học 5 lần qua toàn bộ tập train. Mỗi lần được gọi là một epoch."],
        ["6", "Sau mỗi epoch, kiểm tra trên tập validation và lưu phiên bản có kết quả tốt nhất."],
        ["7", "Dùng phiên bản đã chọn để chạy một lần trên tập test 1.759 mẫu."],
    ], [0.8, 6.0], 11)
    add_figure(doc, DIAGRAMS / "4_ai_pipeline.png", "Hình 3.2. Quy trình huấn luyện và dự đoán", 6.2)

    add_heading(doc, "5. Cấu hình train")
    add_table(doc, ["Thông số", "Giá trị", "Ý nghĩa"], [
        ["Learning rate", "2e-5", "Mức thay đổi của mô hình sau mỗi bước học."],
        ["Batch train", "8", "Mỗi lượt GPU xử lý 8 mẫu."],
        ["Gradient accumulation", "2", "Gộp kết quả của 2 lượt rồi mới cập nhật mô hình."],
        ["Số epoch", "5", "Mô hình đọc toàn bộ tập train 5 lần."],
        ["Weight decay", "0,01", "Hạn chế mô hình ghi nhớ quá sát dữ liệu train."],
        ["Kiểu tính toán", "FP16", "Giúp train nhanh hơn và dùng ít bộ nhớ GPU hơn."],
        ["Cách chọn mô hình", "Macro F1", "Ưu tiên mô hình làm tương đối tốt ở cả ba nhãn."],
    ], [2.0, 1.4, 3.4], 10.5)

    add_heading(doc, "6. Kết quả trong lúc train")
    add_body(doc, "Accuracy là tỷ lệ mẫu được đoán đúng. Macro F1 cho biết mô hình làm tốt đồng đều đến đâu ở cả ba nhãn. Báo cáo dùng cả hai chỉ số vì số mẫu No Fit nhiều hơn hai nhãn còn lại.")
    add_table(doc, ["Epoch", "Train loss", "Val loss", "Accuracy", "Macro F1"], [
        ["1", "1,9763", "0,8560", "53,21%", "0,5374"],
        ["2", "1,5917", "0,7766", "58,44%", "0,5862"],
        ["3", "1,4793", "0,7236", "58,44%", "0,5880"],
        ["4", "1,3009", "0,6924", "65,81%", "0,6576"],
        ["5", "1,1746", "0,6838", "66,03%", "0,6609"],
    ], [1.0, 1.4, 1.4, 1.4, 1.4], 10.5)
    add_figure(doc, ASSETS / "training_metrics.png", "Hình 3.3. Accuracy và Macro F1 qua 5 epoch", 5.6)
    add_body(doc, "Kết quả trên validation tăng dần và tốt nhất ở epoch 5. Tại đây, mô hình đoán đúng 66,03% số mẫu validation và đạt Macro F1 là 0,6609. Phiên bản ở epoch 5 được lưu để chạy bài kiểm tra cuối trên tập test.")

    add_heading(doc, "7. Đã kiểm tra những nội dung nào")
    add_table(doc, ["Nội dung kiểm tra", "Cách kiểm tra", "Mục đích"], [
        ["Đoán nhãn trên dữ liệu mới", "Cho mô hình đoán 1.759 cặp CV-JD trong tập test rồi so với đáp án có sẵn.", "Biết tỷ lệ đoán đúng sau khi train."],
        ["Kết quả của từng nhãn", "Tính F1 riêng cho No Fit, Potential Fit và Good Fit.", "Biết mô hình đang yếu ở nhãn nào."],
        ["So sánh với cách đơn giản", "Chạy TF-IDF và Logistic Regression trên cùng tập test.", "Kiểm tra DistilBERT có tốt hơn cách cơ bản hay không."],
        ["Các trường hợp cụ thể", "Tự tạo ba trường hợp CV và JD dễ hiểu rồi xem nhãn và điểm mô hình trả về.", "Xem mô hình có bị ảnh hưởng quá nhiều bởi từ khóa hay không."],
        ["Khả năng xếp hạng", "Tính NDCG ở 5 và 10 hồ sơ đầu.", "Kiểm tra khả năng sắp xếp hồ sơ phù hợp lên trước."],
        ["Kiểm tra lại 100 mẫu", "Lấy 100 mẫu gần cân bằng giữa ba nhãn và đối chiếu notebook với Excel.", "Kiểm tra thêm kết quả và phát hiện chênh lệch giữa các tệp."],
    ], [2.0, 3.0, 2.0], 9.5)

    add_heading(doc, "8. Kết quả trên 1.759 mẫu test")
    add_body(doc, "TF-IDF kết hợp Logistic Regression là cách so sánh cơ bản. Cách này chủ yếu dựa vào các từ xuất hiện trong CV và JD, sau đó dự đoán một trong ba nhãn.")
    add_table(doc, ["Mô hình", "Accuracy", "Macro F1", "F1 No Fit", "F1 Potential", "F1 Good Fit"], [
        ["TF-IDF + Logistic Regression", "45,59%", "0,4266", "0,5785", "0,3515", "0,3497"],
        ["Mô hình tham chiếu", "58,40%", "0,4320", "0,6900", "0,2800", "0,3260"],
        ["DistilBERT đã train", "44,86%", "0,3779", "0,6015", "0,2099", "0,3224"],
    ], [2.4, 1.0, 1.0, 1.0, 1.1, 1.0], 9.5)
    add_figure(doc, ASSETS / "benchmark_test.png", "Hình 3.4. So sánh kết quả trên tập test", 5.9)
    add_figure(doc, ASSETS / "confusion_matrix_test.png", "Hình 3.5. Ma trận nhầm lẫn trên tập test", 6.15)
    add_body(doc, "DistilBERT đoán đúng 44,86% trong 1.759 mẫu test. Kết quả này thấp hơn TF-IDF khoảng 0,73 điểm phần trăm. Mô hình nhận biết No Fit tốt nhất, nhưng còn nhầm nhiều ở Potential Fit. F1 của Potential Fit chỉ đạt 0,2099.")
    add_body(doc, "Kết quả validation là 66,03% nhưng kết quả test chỉ còn 44,86%. Điều này cho thấy mô hình làm khá tốt trên phần dữ liệu dùng để chọn mô hình, nhưng chưa ổn định khi gặp phần dữ liệu mới được giữ riêng.")

    add_heading(doc, "9. Kiểm tra bằng các trường hợp cụ thể")
    add_table(doc, ["Trường hợp", "Điều muốn kiểm tra", "Kết quả", "Nhận xét"], [
        ["CV Python Backend", "CV có kỹ năng gần với JD.", "Good Fit, 71,1 điểm", "Kết quả phù hợp với nội dung mẫu."],
        ["CV Civil Engineer", "CV khác lĩnh vực với JD công nghệ.", "No Fit, 27,5 điểm", "Mô hình nhận ra trường hợp không phù hợp."],
        ["Junior Data Analyst", "CV có Python và SQL nhưng kinh nghiệm còn ít.", "Good Fit, 82,4 điểm", "Điểm quá cao, cho thấy mô hình dễ bị ảnh hưởng bởi từ khóa."],
    ], [1.8, 2.0, 1.6, 2.0], 9.5)
    add_body(doc, "Phần xếp hạng đạt NDCG@5 là 0,5842 và NDCG@10 là 0,7256. NDCG càng gần 1 thì cách sắp xếp càng tốt. Tuy nhiên, notebook đang tạo nhóm công việc ngẫu nhiên để tính chỉ số này, chưa dùng các ứng viên thật cùng nộp vào một tin tuyển dụng. Vì vậy, kết quả chỉ dùng để tham khảo ban đầu.")

    add_heading(doc, "10. Kiểm tra lại bằng 100 mẫu")
    add_body(doc, "Bài kiểm tra này dùng 100 mẫu gồm 34 No Fit, 33 Potential Fit và 33 Good Fit. Mục đích là xem nhanh kết quả khi số mẫu của ba nhãn gần bằng nhau.")
    add_body(doc, "Notebook ghi nhận 31 mẫu được đoán đúng, tương ứng 31%. Tuy nhiên, file Excel hiện có lại ghi nhận 39 dòng đúng. Hai tệp có thể được tạo từ hai lần chạy khác nhau. Báo cáo tạm dùng con số 31% trong notebook vì đây là nơi lưu toàn bộ quá trình chạy, đồng thời ghi lại chênh lệch để xuất lại file Excel ở lần tiếp theo.")
    add_figure(doc, ASSETS / "test_100_summary.png", "Hình 3.6. Đối chiếu kết quả 100 mẫu", 5.8)

    add_heading(doc, "11. Những điểm cần làm tiếp")
    add_table(doc, ["Vấn đề hiện tại", "Cách xử lý tiếp theo"], [
        ["Kết quả test thấp hơn nhiều so với validation", "Kiểm tra lại cách chia dữ liệu và tránh để cùng một CV hoặc cùng một JD xuất hiện ở nhiều tập."],
        ["Mô hình còn nhầm nhiều ở Potential Fit", "Xem lại cách gán nhãn và bổ sung các mẫu nằm giữa phù hợp và không phù hợp."],
        ["Điểm dễ tăng cao khi có từ khóa giống nhau", "Bổ sung mẫu có cùng kỹ năng nhưng khác số năm kinh nghiệm hoặc khác lĩnh vực."],
        ["CV và JD dài bị cắt bớt", "Chia văn bản thành từng phần để mô hình đọc được nhiều thông tin hơn."],
        ["BERT NER chưa được train", "Train và kiểm tra riêng phần lấy kỹ năng, học vấn và kinh nghiệm từ CV."],
        ["Notebook và Excel chưa trùng kết quả", "Chạy lại một lần, cố định checkpoint và xuất Excel ngay trong lần chạy đó."],
    ], [3.0, 3.8], 10.5)

    add_heading(doc, "12. Kết luận")
    add_body(doc, "Mô hình DistilBERT đã được train đủ 5 epoch trên Tesla T4. Phiên bản tốt nhất đoán đúng 66,03% trên validation, nhưng chỉ đoán đúng 44,86% trên tập test 1.759 mẫu chưa dùng để train.")
    add_body(doc, "Kết quả này cho thấy quy trình train và kiểm tra đã chạy được, nhưng mô hình chưa đủ ổn định để tự động quyết định ứng viên đạt hay không đạt. Trong website, kết quả AI chỉ nên dùng để sắp xếp hồ sơ cho nhà tuyển dụng tham khảo. Nhà tuyển dụng vẫn cần mở CV gốc và tự đưa ra quyết định.")

    add_heading(doc, "13. Tệp dùng để kiểm tra lại kết quả")
    add_table(doc, ["Tệp hoặc nguồn", "Nội dung"], [
        ["kqtrain/TRAIN_CV_JD_FIT_MODEL.ipynb", "Notebook train và toàn bộ kết quả chạy."],
        ["filebaocao/DATASET_TEST_100_SAMPLES_RESULTS.xlsx", "Kết quả kiểm tra 100 mẫu được xuất ra Excel."],
        ["Hugging Face dataset", "cnamuangtoun/resume-job-description-fit."],
        ["Hugging Face model", "distilbert-base-uncased."],
    ], [3.4, 3.4], 10.5)
    doc.save(REPORT3)


if __name__ == "__main__":
    generate_assets()
    build_report2()
    build_report3()
    print(REPORT2)
    print(REPORT3)
